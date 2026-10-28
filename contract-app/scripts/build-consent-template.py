from pathlib import Path
from pypdf import PdfReader
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
source=PdfReader(ROOT/'templates/consent.pdf').pages[0].extract_text()
# Preserve the wording from the source; collapse PDF line breaks and justification spaces only.
start=source.index('в соответствии с Федеральным законом')
end=source.index('Подпись субъекта (представителя субъекта)')
body=source[start:end]
starts=[
 'в соответствии с Федеральным законом',
 'Настоящее согласие предоставляется',
 'Согласие дается на обработку следующих',
 'фамилия,',
 'Согласие  дается  на  обработку',
 'Выражаю ',
 '(согласие/несогласие)',
 'Согласие на обработку и передачу',
 'Субъект персональных данных уведомляется'
]
# The positions are extracted from the PDF, not a retyped version of its legal text.
positions=[]
for prefix in starts:
    if prefix=='(согласие/несогласие)':
        pos=body.index(prefix)
    else:
        pos=body.index(prefix)
    positions.append(pos)
assert positions==sorted(positions)
paragraphs=[' '.join(body[a:b].split()) for a,b in zip(positions,positions[1:]+[len(body)])]

def font(run,size=11,bold=False,italic=False):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.bold=bold;run.italic=italic
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Times New Roman')
    run._element.rPr.rFonts.set(qn('w:cs'),'Times New Roman')
    run.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)

doc=Document();sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=Mm(16);sec.bottom_margin=Mm(16);sec.left_margin=Mm(18);sec.right_margin=Mm(16)
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
normal.paragraph_format.line_spacing=1.0
normal.paragraph_format.space_after=Pt(0)
normal.paragraph_format.widow_control=True
lang=OxmlElement('w:lang');lang.set(qn('w:val'),'ru-RU');normal.element.get_or_add_rPr().append(lang)
doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.title='Согласие на обработку персональных данных';doc.core_properties.comments=''

def add(text,size=11,bold=False,italic=False,align=None,after=0,before=0,keep=False):
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(after);p.paragraph_format.space_before=Pt(before)
    p.paragraph_format.keep_with_next=keep
    if align is not None:p.alignment=align
    font(p.add_run(text),size,bold,italic)
    return p

add('СОГЛАСИЕ\nна обработку персональных данных',12,True,align=WD_ALIGN_PARAGRAPH.CENTER,after=6,keep=True)
add('Я, {fullName}',after=0,keep=True)
add('(фамилия, имя, отчество)',10,italic=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=4,keep=True)
for text in ['Документ, удостоверяющий личность: Паспорт гражданина РФ','Серия и номер: {passportText}','Дата выдачи: {passportDateText}','Наименование выдавшего органа, код подразделения: {passportAuthority}, {passportCode}','Адрес регистрации по месту жительства: {address}']:
    add(text,after=0)
for index,text in enumerate(paragraphs):
    caption=text=='(согласие/несогласие)'
    p=add(text,10 if caption else 11,italic=caption,align=WD_ALIGN_PARAGRAPH.LEFT if caption else WD_ALIGN_PARAGRAPH.JUSTIFY,after=0,before=4 if index==0 else 0,keep=text.startswith('Выражаю'))
    if not caption:p.paragraph_format.first_line_indent=Mm(7)
    else:p.paragraph_format.left_indent=Mm(27)
# A single kept block gives room for handwriting without a caption drifting onto another page.
add('Подпись субъекта (представителя субъекта) персональных данных:',before=8,after=2,keep=True)
for _ in range(2):
    p=add('_'*83,before=6,after=0,keep=True)
    p.paragraph_format.line_spacing=Pt(14)
add('(фамилия, имя, отчество (при наличии) полностью, подпись)',10,italic=True,align=WD_ALIGN_PARAGRAPH.CENTER,after=0,keep=True)
add('«_______» _______________________ ________ г.',before=10,after=0)
doc.save(ROOT/'templates/consent.docx')
print('Created consent.docx from source PDF wording;',len(paragraphs),'body paragraphs')
