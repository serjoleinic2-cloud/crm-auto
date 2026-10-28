"""Verify that the editable template retains the source PDF's body text."""
from pathlib import Path
from pypdf import PdfReader
from docx import Document

root=Path(__file__).resolve().parents[1]
source=PdfReader(root/'templates/consent.pdf').pages[0].extract_text()
text=' '.join(p.text for p in Document(root/'templates/consent.docx').paragraphs)
start='в соответствии с Федеральным законом'
end='Подпись субъекта (представителя субъекта)'
normalize=lambda value:' '.join(value.split())
assert normalize(source[source.index(start):source.index(end)])==normalize(text[text.index(start):text.index(end)])
assert '2026' not in text
assert '«_______» _______________________ ________ г.' in text
footer=text[text.index(end):]
assert '{' not in footer and 'полностью, подпись' in footer
for field in ['fullName','passportText','passportDateText','passportAuthority','passportCode','address']:
    assert text.count('{'+field+'}')==1
print('PASS source PDF wording, top-only placeholders, blank handwritten signature and date')
